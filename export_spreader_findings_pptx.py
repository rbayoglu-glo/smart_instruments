from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE = Path(__file__).parent
OUT = BASE / "output"

PPTX_PATH = OUT / "spreader_findings_slides_2026-08-07.pptx"

FIG_MAIN = OUT / "spreader_theta_at_moment.png"
FIG_THR_THETA = OUT / "engagement_threshold_sensitivity_theta5.png"
FIG_THR_RANK = OUT / "engagement_threshold_sensitivity_rankcorr.png"

CSV_THR = OUT / "engagement_threshold_sensitivity_rankcorr.csv"
CSV_TRIM = OUT / "cycle_trimming_rankcorr_comparison.csv"
CSV_PRE = OUT / "theta0_force_cycle_diagnostics.csv"


TITLE_FONT = Pt(34)
SUBTITLE_FONT = Pt(18)
BODY_FONT = Pt(20)
SMALL_FONT = Pt(14)


def set_title(shape, text: str, size=TITLE_FONT):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = True


def add_bullets(slide, lines):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.8), Inches(5.2))
    tf = box.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = BODY_FONT


def add_table_text(slide, title: str, lines):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12.0), Inches(0.6))
    set_title(title_box, title, size=Pt(26))

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.8), Inches(5.8))
    tf = box.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = SMALL_FONT


def main() -> None:
    thr = pd.read_csv(CSV_THR)
    trim = pd.read_csv(CSV_TRIM)
    preload = (
        pd.read_csv(CSV_PRE)
        .groupby("intervention", as_index=False)["F_start_N"]
        .median()
        .sort_values("F_start_N", ascending=False)
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: title
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_title(s.shapes.title, "Lamina Spreader Findings")
    subtitle = s.placeholders[1].text_frame
    subtitle.clear()
    p = subtitle.paragraphs[0]
    p.text = "L231147 L5-S1 | 2026-08-07"
    p.font.size = SUBTITLE_FONT

    # Slide 2: scope and decisions
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s.shapes.title, "Scope And Final Decisions", size=Pt(28))
    add_bullets(
        s,
        [
            "Included: Intact Holding, PUBF Left, FUF Left, FBF, Posterior Release, SPO",
            "Excluded: baseline Intact",
            "Primary metric: median theta at fixed pivot moments (2, 3, 4, 5 N.m)",
            "Primary reference: theta from tip engagement",
            "Engagement threshold selected: 12 N",
        ],
    )

    # Slide 3: why force near theta=0
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s.shapes.title, "Why Force Near Theta = 0 Can Be High", size=Pt(28))
    lines = [
        "Force zero and theta zero are referenced to different events.",
        "Preload and seating can occur before meaningful opening.",
        "Median ramp-start preload by intervention:",
    ]
    for _, r in preload.iterrows():
        lines.append(f"  - {r['intervention']:<18} {r['F_start_N']:.3f} N")
    add_bullets(s, lines)

    # Slide 4: primary figure
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s.shapes.title, "Primary Comparator: Theta From Tip Engagement", size=Pt(28))
    s.shapes.add_picture(str(FIG_MAIN), Inches(0.6), Inches(0.9), width=Inches(12.1))

    # Slide 5: threshold sensitivity table
    s = prs.slides.add_slide(prs.slide_layouts[5])
    table_lines = [
        "Threshold  rho@2   rho@3   rho@4   rho@5",
        "-----------------------------------------",
    ]
    for _, r in thr.iterrows():
        table_lines.append(
            f"{r['threshold_N']:>8.1f}  {r['rank_corr@2.0']:>+6.3f}  {r['rank_corr@3.0']:>+6.3f}  "
            f"{r['rank_corr@4.0']:>+6.3f}  {r['rank_corr@5.0']:>+6.3f}"
        )
    table_lines.append("")
    table_lines.append("Best overall balance across moments: 12 N")
    add_table_text(s, "Engagement Threshold Sensitivity", table_lines)

    # Slide 6: threshold sensitivity figures
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s.shapes.title, "Threshold Sensitivity Figures", size=Pt(28))
    s.shapes.add_picture(str(FIG_THR_THETA), Inches(0.5), Inches(1.0), width=Inches(6.2))
    s.shapes.add_picture(str(FIG_THR_RANK), Inches(6.7), Inches(1.0), width=Inches(6.2))

    # Slide 7: cycle trimming impact
    s = prs.slides.add_slide(prs.slide_layouts[5])
    trim_lines = [
        "Drop first 3 and last cycle per intervention:",
        "",
        "Moment   all_rho   trimmed_rho   delta",
        "--------------------------------------",
    ]
    for _, r in trim.iterrows():
        trim_lines.append(
            f" {r['moment_Nm']:>4.1f}    {r['rank_corr_all']:>+7.3f}      {r['rank_corr_trim']:>+7.3f}    {r['delta_rank_corr']:>+7.3f}"
        )
    trim_lines += [
        "",
        "Conclusion: blanket cycle trimming reduced stability and coverage.",
    ]
    add_table_text(s, "Cycle Trimming Impact", trim_lines)

    # Slide 8: recommendations
    s = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s.shapes.title, "Recommendations", size=Pt(28))
    add_bullets(
        s,
        [
            "1) Keep tip-engagement reference as the primary comparator.",
            "2) Keep engagement threshold at 12 N for this dataset.",
            "3) Do not remove first 3 and last cycle by default.",
            "4) If more filtering is needed, use quality-based cycle rules.",
        ],
    )

    prs.save(PPTX_PATH)
    print(f"saved {PPTX_PATH}")


if __name__ == "__main__":
    main()
