import argparse
import csv
from datetime import datetime, time, timedelta, timezone
import hashlib
import json
from pathlib import Path
import re

import openpyxl

from intact_level_processing_config import DATA_ROOT, OUTPUT_DIR


WORKBOOK_PATH = DATA_ROOT / "segment" / "spreader" / "displacement_and_force_correspondence.xlsx"
OUTPUT_PATH = OUTPUT_DIR / "intraoperative_measurement_file_audit.csv"
ALL_OUTPUT_PATH = OUTPUT_DIR / "spreader_measurement_file_audit.csv"
SPREADER_ROOT = DATA_ROOT / "segment" / "spreader"
TRACKING_ROOT = SPREADER_ROOT / "Displacement" / "gns-sensorhub"
LUMBAR_ACQUISITION_DATE = "2026-09-04"
INTRAOP_SHEETS = ("Thoracic, Prone Specimen", "Lumbar, Prone Specimen")
FORCE_FACTS = {}


def text(value) -> str:
    return "" if value is None else str(value).strip()


def normalize_level(value: str) -> str:
    match = re.fullmatch(r"([TL])(\d+)\s*-\s*([TL]?)(\d+)", value.strip().upper())
    if match is None:
        raise ValueError(f"Unrecognized workbook level: {value!r}")
    first_region, first_number, second_region, second_number = match.groups()
    return f"{first_region}{first_number}-{second_region or first_region}{second_number}"


def force_facts(path: Path) -> dict:
    if path in FORCE_FACTS:
        return FORCE_FACTS[path]
    facts = {"exists": path.is_file(), "duration_s": "", "sha256": ""}
    if path.is_file():
        facts["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            first_time = None
            last_time = None
            for record in reader:
                sample_time = float(record["Time"])
                if first_time is None:
                    first_time = sample_time
                last_time = sample_time
            if first_time is not None:
                facts["duration_s"] = round((last_time - first_time) / 1e6, 3)
    FORCE_FACTS[path] = facts
    return facts


def recording_start(folder: Path) -> datetime:
    return datetime.strptime(folder.name, "%Y-%m-%d_[%H-%M-%S]")


def clock(value) -> time:
    if isinstance(value, time):
        return value
    return time.fromisoformat(text(value))


def tracking_catalog() -> list[dict]:
    recordings = {}
    pending = [TRACKING_ROOT]
    while pending:
        parent = pending.pop()
        for folder in sorted(parent.iterdir()):
            if not folder.is_dir() or folder.name.startswith("2026-11-"):
                continue
            metadata_path = folder / "recordingMetadata.json"
            if not metadata_path.is_file():
                if re.fullmatch(r"[0-9a-fA-F-]{36}", folder.name):
                    pending.append(folder)
                continue
            with metadata_path.open(encoding="utf-8-sig") as stream:
                metadata = json.load(stream)
            start = metadata.get("start", {})
            date_time = text(start.get("date_time"))
            if not date_time or date_time.startswith("2026-11-"):
                continue
            utc_start = datetime.fromisoformat(date_time)
            recorded_start = utc_start.astimezone(timezone(timedelta(hours=-5)))
            identity = (date_time, text(start.get("timestamp")))
            recording = {
                "folder": folder,
                "date": recorded_start.date().isoformat(),
                "start": recorded_start,
            }
            if identity not in recordings or len(folder.parts) < len(recordings[identity]["folder"].parts):
                recordings[identity] = recording
    return sorted(recordings.values(), key=lambda recording: str(recording["folder"]))


def tracking_duration(folder: Path) -> float | None:
    pose_path = next(folder.glob("ndi/*/NDI_Localizer.txt"), folder / "pf" / "poses.txt")
    if not pose_path.is_file():
        return None
    first_timestamp = None
    last_timestamp = None
    with pose_path.open(encoding="utf-8-sig") as stream:
        for fields in csv.reader(stream):
            if not fields:
                continue
            try:
                timestamp = float(fields[0])
            except ValueError:
                continue
            if first_timestamp is None:
                first_timestamp = timestamp
            last_timestamp = timestamp
    if first_timestamp is None or last_timestamp is None:
        return None
    return (last_timestamp - first_timestamp) / 1e6


def lumbar_tracking(record: dict, recordings: list[dict], force_duration) -> tuple[Path | None, str, str, str]:
    workbook_clock = clock(record["workbook_computer_time"])
    clock_seconds = workbook_clock.hour * 3600 + workbook_clock.minute * 60 + workbook_clock.second
    candidates = []
    for recording in recordings:
        recorded_clock = recording["start"].time()
        recorded_seconds = recorded_clock.hour * 3600 + recorded_clock.minute * 60 + recorded_clock.second
        difference = min(
            abs(recorded_seconds - clock_seconds),
            abs(recorded_seconds - (clock_seconds + 12 * 3600)),
        )
        if difference > 120:
            continue
        duration = tracking_duration(recording["folder"])
        duration_ok = duration is not None and force_duration != "" and abs(duration - force_duration) <= max(10.0, 0.10 * force_duration)
        candidates.append({
            "tracking_folder": str(recording["folder"]),
            "recorded_date": recording["date"],
            "clock_difference_s": round(difference, 3),
            "tracking_duration_s": None if duration is None else round(duration, 3),
            "duration_matches": duration_ok,
            "acquisition_date_matches": recording["date"] == LUMBAR_ACQUISITION_DATE,
        })
    verified = [candidate for candidate in candidates if candidate["duration_matches"] and candidate["acquisition_date_matches"]]
    if len(verified) == 1:
        candidate = verified[0]
        return Path(candidate["tracking_folder"]), "workbook_clock_duration_acquisition_date_match", str(candidate["clock_difference_s"]), json.dumps(candidates)
    status = "multiple_verified_candidates_unresolved" if verified else "no_non_november_acquisition_date_clock_duration_match"
    return None, status, "", json.dumps(candidates)


def worksheet_records(sheet) -> list[dict]:
    header_row = None
    headers = {}
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        candidate_headers = {text(value).casefold(): column for column, value in enumerate(values) if value is not None}
        if "level" in candidate_headers and "session number" in candidate_headers:
            header_row = row_number
            headers = candidate_headers
            break
    if header_row is None:
        raise ValueError(f"Required headers were not found in {sheet.title!r}.")

    records = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        level_value = text(values[headers["level"]])
        session = text(values[headers["session number"]])
        if not level_value or not session:
            continue
        if "lamina spreader side" in headers:
            level = normalize_level(level_value)
            side = text(values[headers["lamina spreader side"]]).title()
            computer_time = values[headers["computer time start"]]
            other_time = values[headers["est spreading start"]]
        else:
            level_text, side = level_value.rsplit(None, 1)
            level = normalize_level(level_text)
            side = side.title()
            computer_time = values[headers["computer clock"]]
            other_time = values[headers["est"]]
        if side not in {"Left", "Right"}:
            raise ValueError(f"Invalid side at {sheet.title}:{row_number}: {side!r}")
        notes = text(values[headers["notes"]]) if "notes" in headers else ""
        records.append({
            "worksheet": sheet.title,
            "worksheet_row": row_number,
            "spine_level": level,
            "side": side,
            "workbook_session": session,
            "workbook_computer_time": clock(computer_time).isoformat(),
            "workbook_other_time": clock(other_time).isoformat(),
            "intervention": text(values[headers["intervention"]]) if "intervention" in headers else "",
            "workbook_notes": notes,
        })
    return records


def thoracic_tracking(sheet, record: dict) -> tuple[Path | None, str, str]:
    date_match = re.search(r"\d{4}-\d{2}-\d{2}", text(sheet["A2"].value))
    if date_match is None:
        return None, "workbook_recording_date_missing", ""
    recording_date = datetime.strptime(date_match.group(), "%Y-%m-%d").date()
    workbook_start = datetime.combine(recording_date, clock(record["workbook_computer_time"]))
    tracking_root = SPREADER_ROOT / "Displacement" / "gns-sensorhub"
    candidates = []
    for folder in tracking_root.glob(f"{recording_date.isoformat()}_*"):
        if not folder.is_dir() or re.fullmatch(r"\d{4}-\d{2}-\d{2}_\[\d{2}-\d{2}-\d{2}\]", folder.name) is None:
            continue
        difference = (recording_start(folder) - workbook_start).total_seconds()
        if abs(difference) <= 120:
            candidates.append((folder, difference))
    if len(candidates) != 1:
        status = "no_recording_within_two_minutes" if not candidates else "multiple_time_candidates_unresolved"
        return None, status, ""
    folder, difference = candidates[0]
    metadata_path = folder / "recordingMetadata.json"
    if not metadata_path.is_file():
        return folder, "time_candidate_metadata_missing", str(difference)
    with metadata_path.open(encoding="utf-8-sig") as stream:
        metadata = json.load(stream)
    if "date_time" not in metadata.get("start", {}):
        return folder, "time_candidate_metadata_incomplete", str(difference)
    return folder, "workbook_clock_candidate_not_signal_verified", str(difference)


def segment_reference(workbook, force_file: str) -> str:
    session_match = re.fullmatch(r"Session #(\d+)\.csv", force_file)
    if session_match is None:
        return ""
    session_number = session_match.group(1)
    for sheet in workbook:
        if sheet.title in INTRAOP_SHEETS:
            continue
        session_column = None
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if session_column is None:
                headers = {text(value).casefold(): column for column, value in enumerate(values) if value is not None}
                session_column = headers.get("session number")
                continue
            if text(values[session_column]) == session_number:
                return f"{sheet.title}!row {row_number}"
    return ""


def build_audit(workbook) -> list[dict]:
    recordings = tracking_catalog()
    previous_records = {}
    if OUTPUT_PATH.is_file():
        with OUTPUT_PATH.open(encoding="utf-8-sig", newline="") as stream:
            previous_records = {
                (record["worksheet"], record["spine_level"], record["side"]): record
                for record in csv.DictReader(stream)
            }
    records = []
    for sheet_name in INTRAOP_SHEETS:
        sheet = workbook[sheet_name]
        for source in worksheet_records(sheet):
            record = {
                "specimen": DATA_ROOT.name,
                "measurement_type": "intraoperative",
                "acquisition_date": LUMBAR_ACQUISITION_DATE if sheet_name == "Lumbar, Prone Specimen" else "",
                "acquisition_date_evidence": "user_confirmed" if sheet_name == "Lumbar, Prone Specimen" else "",
                **source,
            }
            trial_name = f"{record['spine_level']} {record['side']}"
            flags = []
            candidate_evidence = ""
            if sheet_name == "Thoracic, Prone Specimen":
                force_path = SPREADER_ROOT / "Force" / "lamspreader" / f"Session #{record['workbook_session']}.csv"
                tracking_folder, tracking_status, time_difference = thoracic_tracking(sheet, record)
                if tracking_folder is None:
                    flags.append("tracking_assignment_unresolved")
            else:
                force_path = SPREADER_ROOT / "Force" / "lamspreader" / f"{trial_name}.csv"
                tracking_folder, tracking_status, time_difference, candidate_evidence = lumbar_tracking(
                    record, recordings, force_facts(force_path)["duration_s"]
                )
                if tracking_folder is None:
                    flags.append("cadaver_lumbar_tracking_pair_not_verified")
            facts = force_facts(force_path)
            record.update({
                "source_force_csv": str(force_path),
                "source_force_exists": facts["exists"],
                "source_force_duration_s": facts["duration_s"],
                "source_force_sha256": facts["sha256"],
                "source_tracking_folder": str(tracking_folder) if tracking_folder else "",
                "source_tracking_exists": tracking_folder.is_dir() if tracking_folder else False,
                "tracking_mapping_status": tracking_status,
                "tracking_minus_workbook_start_s": time_difference,
                "non_november_candidate_evidence": candidate_evidence,
                "previous_analysis_status": "not_used_in_previous_lumbar_step_charts",
                "previous_force_csv": "",
                "previous_tracking_folder": "",
                "previous_measurement_type": "",
                "previous_workbook_reference": "",
                "previous_force_sha256": "",
            })
            previous = previous_records.get((sheet_name, record["spine_level"], record["side"]), {})
            if previous.get("previous_force_csv"):
                reference = segment_reference(workbook, Path(previous["previous_force_csv"]).name)
                used_tracking = previous.get("previous_tracking_folder", "")
                excluded_tracking = "2026-11-" in used_tracking
                record.update({
                    "previous_analysis_status": "excluded_specimen_tracking_used" if excluded_tracking else previous.get("previous_analysis_status", "requires_review"),
                    "previous_force_csv": previous["previous_force_csv"],
                    "previous_tracking_folder": used_tracking,
                    "previous_measurement_type": previous.get("previous_measurement_type", "unresolved"),
                    "previous_workbook_reference": reference or previous.get("previous_workbook_reference", ""),
                    "previous_force_sha256": previous.get("previous_force_sha256", ""),
                })
                if reference:
                    flags.append("segment_measurement_used_in_previous_intraoperative_comparison")
                if excluded_tracking:
                    flags.append("excluded_november_specimen_tracking_used_in_previous_analysis")
            if "swapped" in record["workbook_notes"].casefold():
                flags.append("possible_side_swap_per_workbook_not_corrected")
            if record["workbook_notes"]:
                flags.append("see_original_workbook_note")
            if not facts["exists"]:
                flags.append("force_file_missing")
            record["audit_flags"] = "; ".join(flags)
            records.append(record)
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit cadaver spreader file provenance.")
    parser.add_argument("--force-review", action="store_true")
    args = parser.parse_args()
    workbook_hash = hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        records = build_audit(workbook)
    finally:
        workbook.close()
    for record in records:
        record["workbook_file"] = str(WORKBOOK_PATH)
        record["workbook_sha256"] = workbook_hash
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with OUTPUT_PATH.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    print(f"Saved {len(records)} workbook records to {OUTPUT_PATH}")
    for sheet_name in INTRAOP_SHEETS:
        print(f"{sheet_name}: {sum(record['worksheet'] == sheet_name for record in records)} records")
    for record in records:
        if record["audit_flags"]:
            print(f"{record['worksheet']} row {record['worksheet_row']} | {record['spine_level']} {record['side']} | {record['audit_flags']}")
    listed = {(record["spine_level"], record["side"]) for record in records}
    print(f"T1-T2 Right listed in the workbook: {('T1-T2', 'Right') in listed}; no file assignment inferred.")

    if args.force_review:
        from plot_l4l5_intact_processing_steps import plot_lumbar_force_review
        plot_lumbar_force_review(records, OUTPUT_DIR)


if __name__ == "__main__":
    main()